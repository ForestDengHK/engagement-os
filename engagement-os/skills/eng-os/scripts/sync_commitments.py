#!/usr/bin/env python3
"""Make the drafting follow the commitments register — one place to change a value or approval.

`3_drafting/commitments.md` holds every commitment the response states, once. A bid lead changes
a Value (or an approver answers in the sign-off workbook); this script carries the change into
the sections, the sign-off sheet and the deck, and proves nothing drifted. The judgement part —
a changed value that needs an argument rewritten, not a string swapped — stays with
`eng-propagate-change`; this is only the deterministic part.

Register columns (located by header name, never position):
    | ID | Topic | Class | Value | Applied value | Basis | Approver | Status | Sections | Slides |
`Value` is what the response must say; `Applied value` is what the listed sections say now.

Usage:
    python3 sync_commitments.py <repo-root> [check|apply|signoff|import-xlsx <xlsx>|rebuild|all]
                                [--eng-id X] [--strict]

    check        every listed section states the Applied value (or Value, before the first
                 sync) verbatim in its body; reports drift, unsynced and rejected rows; then
                 deck-drift: every row with Slides states its Value in the newest _render deck
    apply        replaces Applied value → Value in each listed section body (never in the
                 frontmatter or review log), reopens reviewed sections, logs the change
    signoff      regenerates 3_drafting/signoff_sheet.md and signoff_sheet.xlsx (decisions
                 to approve · firm inputs owed · all commitments)
    import-xlsx  reads the approver's answers back into the register, then previews `apply`
    rebuild      deck manifest (draft profile) → _render/rebuild.sh if present → verify_deck
                 → deck-drift
    all          apply → signoff → rebuild → check --strict

Exit: 0 ok · 1 drift / failed step (check: only with --strict) · 2 no register / bad input ·
      3 a register value missing from the deck (deck-drift)
"""
import argparse
import collections
import datetime as dt
import html
import os
import pathlib
import re
import subprocess
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import section_contract as sc

SCRIPTS = pathlib.Path(__file__).resolve().parent
REGISTER_REL = "3_drafting/commitments.md"
STATUSES = ("decided", "pending", "approved", "rejected", "changed", "confirmed")
#: D rows on these topics change a commercial or contractual position → someone must sign them
SIGNOFF_TOPIC_RE = re.compile(r"price|rate|discount|hold|indexation|liability|acceptance|"
                              r"warranty|SLA|lead time|named person", re.I)
REOPEN = {"reviewed-r1", "reviewed-r2", "approved"}
LOG_RE = re.compile(r"^##\s+Review log\s*$", re.M | re.I)
LOG_HEAD = ("| Round | Reviewer / lens | Date | Verdict | What changed |\n"
            "|---|---|---|---|---|\n")
HEAD_FILL = "86BC25"

DECISION_COLS = ["ID", "Topic", "Value", "Sections", "Slides", "Basis", "Approver", "Status"]
FIRM_COLS = ["ID", "Priority", "Topic", "Value (fallback if not confirmed)", "Owner",
             "Sections", "Basis", "Status"]
ALL_COLS = ["ID", "Topic", "Class", "Value", "Applied value", "Basis", "Approver", "Status",
            "Sections", "Slides"]
ALL_KEYS = ["id", "topic", "class", "value", "applied", "basis", "approver", "status",
            "sections", "slides"]
REVIEW_COLS = ["Approved? (Y/N)", "New value", "Comment"]


# ── the register ──────────────────────────────────────────────────────────────

def _key(header):
    h = header.lower().strip("*` ")
    if h == "id":
        return "id"
    if h.startswith("applied"):
        return "applied"
    if h.startswith("value"):
        return "value"
    for key in ("topic", "class", "basis", "approver", "status", "section", "slide"):
        if h.startswith(key):
            return {"section": "sections", "slide": "slides"}.get(key, key)
    return None


def split_cells(line):
    """Pipe-table cells; an escaped `\\|` stays inside its cell."""
    cells = re.split(r"(?<!\\)\|", line.strip())
    return [c.strip() for c in cells[1:-1]]


def _unescape(cell):
    return cell.replace("\\|", "|")


def _escape(value):
    return str(value).replace("|", "\\|").replace("\n", " ")


class Register:
    """The commitments table, editable in place — every other line of the file is kept."""

    def __init__(self, path):
        self.path = pathlib.Path(path)
        self.pursuit = self.path.parents[1]
        self.lines = self.path.read_text(encoding="utf-8").splitlines()
        self.cols, self.rows = None, []
        for i, line in enumerate(self.lines):
            if not line.lstrip().startswith("|"):
                if self.cols and self.rows:
                    break
                continue
            cells = split_cells(line)
            if self.cols is None:
                keys = [_key(c) for c in cells]
                if {"id", "value", "sections"} <= set(keys):
                    self.cols = {k: n for n, k in enumerate(keys) if k and k not in
                                 {kk for kk in keys[:n]}}
                continue
            if re.fullmatch(r"[\s:-]*", "".join(cells)):
                continue                                  # the separator row
            row = {k: _unescape(cells[n]) if n < len(cells) else ""
                   for k, n in self.cols.items()}
            if "<" in row["id"] or re.fullmatch(r"<.*>", row["value"]) or \
                    "<" in row["sections"]:
                continue                                  # planted template row
            row["_line"], row["_cells"] = i, cells
            self.rows.append(row)
        if self.cols is None:
            raise ValueError(f"{self.path}: no table with ID / Value / Sections columns")

    def sections(self, row):
        names = [n.strip().strip("`") for n in re.split(r"[,;]", row["sections"])]
        return [n for n in names if n and n not in ("—", "-")]

    def set(self, row, **changes):
        for key, val in changes.items():
            if key not in self.cols:
                raise ValueError(f"{self.path.name} has no `{key}` column — use the contract "
                                 "columns (ID | Topic | Class | Value | Applied value | Basis | "
                                 "Approver | Status | Sections | Slides)")
            row[key] = val
            cells = row["_cells"]
            while len(cells) <= self.cols[key]:
                cells.append("")
            cells[self.cols[key]] = _escape(val)
        self.lines[row["_line"]] = "| " + " | ".join(row["_cells"]) + " |"

    def save(self):
        self.path.write_text("\n".join(self.lines) + "\n", encoding="utf-8")


def register_paths(root, eng_id=None):
    root = pathlib.Path(root)
    if eng_id:
        p = root / "01_pursuit" / eng_id / REGISTER_REL
        return [p] if p.exists() else []
    return sorted(root.glob(f"01_pursuit/*/{REGISTER_REL}"))


def section_index(pursuit):
    base = pursuit / "3_drafting/sections"
    return {p.name: p for p in sorted(base.rglob("*.md"))} if base.exists() else {}


def resolve(index, name):
    return index.get(name) or index.get(name + ".md")


def body_span(text):
    """(start, end) of the prose a commitment may live in: after the frontmatter, before the
    review log. The frontmatter is metadata and the log is history — neither is rewritten."""
    m = sc.FM_RE.match(text)
    start = m.end() if m else 0
    log = LOG_RE.search(text, start)
    return start, (log.start() if log else len(text))


def body_of(path):
    text = path.read_text(encoding="utf-8", errors="replace")
    s, e = body_span(text)
    return text[s:e]


# ── check ─────────────────────────────────────────────────────────────────────

def check(reg):
    """{'drift': [(id, file, msg)], 'unsynced': [row], 'rejected': [row], 'never': [row],
    'counts': Counter(status)} — the logic eng_lint's commitment rules import."""
    index = section_index(reg.pursuit)
    out = {"drift": [], "unsynced": [], "rejected": [], "never": [],
           "counts": collections.Counter(r.get("status", "").lower() or "—" for r in reg.rows)}
    for row in reg.rows:
        applied, value = row.get("applied", ""), row["value"]
        status = row.get("status", "").lower()
        needle = applied or value
        if needle:
            for name in reg.sections(row):
                path = resolve(index, name)
                if path is None:
                    out["drift"].append((row["id"], name, "section file not found"))
                elif needle not in body_of(path):
                    out["drift"].append((row["id"], name, "applied value not found" if applied
                                         else "value not found"))
        if status == "changed" or (applied and applied != value):
            out["unsynced"].append(row)
        elif not applied and value and reg.sections(row):
            out["never"].append(row)
        if status == "rejected" and (not applied or applied == value):
            out["rejected"].append(row)
    return out


# ── deck drift ────────────────────────────────────────────────────────────────

AT_RE = re.compile(r"<a:t(?:\s[^>]*)?>([^<]*)</a:t>")


def newest_deck(pursuit):
    render = pursuit / "3_drafting/_render"
    decks = [d for d in render.rglob("*.pptx") if not d.name.startswith("~$")] \
        if render.exists() else []
    return max(decks, key=lambda d: d.stat().st_mtime) if decks else None


def deck_text(pptx):
    """All slide text, runs joined within a paragraph (PowerPoint splits a phrase across runs),
    whitespace collapsed. stdlib zipfile + the `<a:t>` element — no python-pptx."""
    parts = []
    with zipfile.ZipFile(pptx) as z:
        for name in z.namelist():
            if re.fullmatch(r"ppt/slides/slide\d+\.xml", name):
                xml = z.read(name).decode("utf-8", errors="replace")
                parts += ["".join(AT_RE.findall(para)) for para in xml.split("</a:p>")]
    return " ".join(" ".join(html.unescape(p) for p in parts).split())


def deck_drift(reg):
    """(deck path or None, [(id, value, slides)] rows whose Value is not in the deck)."""
    deck = newest_deck(reg.pursuit)
    if deck is None:
        return None, []
    text = deck_text(deck)
    misses = [(r["id"], r["value"], r.get("slides", "")) for r in reg.rows
              if r.get("slides") and r.get("applied") and
              " ".join(r["value"].split()) not in text]
    return deck, misses


def report_deck_drift(reg, root):
    deck, misses = deck_drift(reg)
    if deck is None:
        print("deck-drift: no .pptx under 3_drafting/_render/ — skipped")
        return 0
    print(f"deck-drift: {deck.relative_to(root)} · {len(misses)} miss(es)")
    for cid, val, slides in misses:
        print(f"  MISS {cid}: '{val}' not in deck (slides {slides})")
    return 3 if misses else 0


def cmd_check(reg, root, strict):
    res = check(reg)
    print(f"check: {reg.path.relative_to(root)} · {len(reg.rows)} row(s)")
    for cid, name, msg in res["drift"]:
        print(f"  drift     {cid} · {name} · {msg}")
    for row in res["unsynced"]:
        print(f"  unsynced  {row['id']} · \"{row.get('applied', '')}\" → \"{row['value']}\" "
              f"(status {row.get('status', '')}) — run apply")
    for row in res["rejected"]:
        print(f"  rejected  {row['id']} · rejected with no new Value — write the new Value "
              "in the register (or the sign-off workbook) and sync")
    if res["never"]:
        print(f"  ({len(res['never'])} row(s) not yet synced — Applied value empty; "
              "apply records it)")
    print("  status: " + " · ".join(f"{k} {v}" for k, v in sorted(res["counts"].items())))
    bad = res["drift"] or res["unsynced"]
    if not bad and not res["rejected"]:
        print("  ✓ no drift, nothing unsynced")
    if report_deck_drift(reg, root):
        return 3
    return 1 if strict and bad else 0


# ── apply ─────────────────────────────────────────────────────────────────────

def add_log_row(text, row):
    """Append `row` to the review-log table — above any row planted for a round not yet run."""
    m = LOG_RE.search(text)
    if not m:
        return text.rstrip("\n") + "\n\n## Review log\n\n" + LOG_HEAD + row + "\n"
    pos, last = m.end(), None
    for line in text[m.end():].splitlines(keepends=True):
        s = line.strip()
        if s.startswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            vals = [re.sub(r"<[^>]*>", "", c).strip() for c in cells[2:4]]
            if cells and sc.ROUND_LABEL_RE.match(cells[0]) and len(cells) > 3 and \
                    not any(vals):
                break
            last = pos + len(line)
        elif last is not None and s:
            break
        pos += len(line)
    if last is None:
        return text[:m.end()] + "\n\n" + LOG_HEAD + row + "\n" + text[m.end():]
    head = text[:last] if text[:last].endswith("\n") else text[:last] + "\n"
    return head + row + "\n" + text[last:]


def reopen(path, changes, today):
    """Send a reviewed section back to revise-r1 and log each commitment that moved under it.

    draft / blocked sections keep their status and get no log row: a draft has no round to
    reopen, and a `revise` verdict under `blocked-r1` would contradict its status."""
    text = path.read_text(encoding="utf-8")
    meta, _ = sc.parse_frontmatter(text)
    status = (meta.get("status") or "").strip()
    new_status = "revise-r1" if status in REOPEN else status
    if new_status != status:
        text = re.sub(r"^status:\s*\S+.*$", f"status: {new_status}", text, count=1, flags=re.M)
    if not new_status.startswith("revise"):
        path.write_text(text, encoding="utf-8")
        return status, new_status
    label = sc.round_of(new_status)
    for cid, old, new in changes:
        what = f'commitment {cid} changed: "{old}" → "{new}"'.replace("|", "/")
        text = add_log_row(text, f"| {label} (sync) | sync_commitments | {today} | revise | "
                                 f"{what} |")
    path.write_text(text, encoding="utf-8")
    return status, new_status


def apply(reg, root, dry_run=False):
    """Carry every changed Value into its sections. Returns (report lines, drift found)."""
    index = section_index(reg.pursuit)
    verb = "would replace" if dry_run else "replaced"
    lines, drift, per_section = [], False, collections.OrderedDict()
    for row in reg.rows:
        old, new = row.get("applied", ""), row["value"]
        status = row.get("status", "").lower()
        if old == new and status != "changed":
            continue
        if not new:
            continue
        missing, edits = [], []
        for name in reg.sections(row):
            path = resolve(index, name)
            if path is None:
                missing.append((name, "section file not found"))
                continue
            text = path.read_text(encoding="utf-8")
            s, e = body_span(text)
            body = text[s:e]
            if old and old != new and old in body:
                for m in re.finditer(re.escape(old), body):
                    ln = text.count("\n", 0, s + m.start()) + 1
                    lines.append(f"  {verb} {row['id']} · {path.name}:{ln} · \"{old}\" → \"{new}\"")
                edits.append((path, text[:s] + body.replace(old, new) + text[e:]))
            elif new in body:
                continue                                  # already says the new value
            else:
                missing.append((name, "applied value not found" if old else "value not found"))
        for name, msg in missing:
            lines.append(f"  drift   {row['id']} · {name} · {msg}")
        if missing:
            drift = True
            continue
        if not old:
            lines.append(f"  {'would record' if dry_run else 'recorded'} {row['id']} · "
                         "Applied value (first sync)")
        if dry_run:
            continue
        for path, text in edits:
            path.write_text(text, encoding="utf-8")
            per_section.setdefault(path, []).append((row["id"], old, new))
        upd = {"applied": new}
        if status in ("changed", "rejected"):
            upd["status"] = "decided"
        reg.set(row, **upd)
    if not dry_run:
        today = dt.date.today().isoformat()
        for path, changes in per_section.items():
            before, after = reopen(path, changes, today)
            note = f"{before} → {after}" if before != after else f"{after} (unchanged)"
            lines.append(f"  section {path.name} · status {note}")
        reg.save()
    return lines, drift


def cmd_apply(reg, root):
    lines, drift = apply(reg, root)
    print(f"apply: {reg.path.relative_to(root)}")
    print("\n".join(lines) if lines else "  nothing to apply — every Applied value is current")
    return 1 if drift else 0


# ── signoff ───────────────────────────────────────────────────────────────────

def _md_tables(path):
    """[(headers, [cells])] for every pipe table in `path`."""
    out, cur = [], None
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.lstrip().startswith("|"):
            cells = [_unescape(c) for c in split_cells(line)]
            if cur is None:
                cur = (cells, [])
                out.append(cur)
            elif not re.fullmatch(r"[\s:-]*", "".join(cells)):
                cur[1].append(cells)
        else:
            cur = None
    return out


def signoff_rows(reg, old_md=None):
    prio = {}
    if old_md and old_md.exists():
        for headers, rows in _md_tables(old_md):
            if "ID" in headers and "Priority" in headers:
                i, p = headers.index("ID"), headers.index("Priority")
                prio.update({r[i]: r[p] for r in rows if len(r) > max(i, p)})
    decisions, firm = [], []
    for row in reg.rows:
        cls = row.get("class", "").strip().upper()[:1]
        if cls == "D" and SIGNOFF_TOPIC_RE.search(row.get("topic", "")):
            decisions.append([row["id"], row.get("topic", ""), row["value"], row["sections"],
                              row.get("slides", ""), row.get("basis", ""),
                              row.get("approver", ""), row.get("status", "")])
        elif cls == "X":
            firm.append([row["id"], prio.get(row["id"], "2"), row.get("topic", ""),
                         row["value"], row.get("approver", ""), row["sections"],
                         row.get("basis", ""), row.get("status", "")])
    every = [[row.get(k, "") for k in ALL_KEYS] for row in reg.rows]
    return decisions, firm, every


def _md_table(cols, rows):
    out = ["| " + " | ".join(cols) + " |", "|" + "---|" * len(cols)]
    out += ["| " + " | ".join(_escape(c) for c in r) + " |" for r in rows]
    return "\n".join(out)


def cmd_signoff(reg, root):
    draft = reg.pursuit / "3_drafting"
    md_path, xlsx_path = draft / "signoff_sheet.md", draft / "signoff_sheet.xlsx"
    decisions, firm, every = signoff_rows(reg, md_path)
    deadline = "**Submission deadline:** <date/time — §cite>"
    if md_path.exists():
        m = re.search(r"^\*\*Submission deadline:\*\*.*$", md_path.read_text(encoding="utf-8"),
                      re.M)
        deadline = m.group(0) if m else deadline
    md_path.write_text(
        f"# Sign-off sheet — {reg.pursuit.name}\n\n"
        "> Generated by `sync_commitments.py signoff` from `commitments.md` — change the "
        "register, not this file, and run `/eng-sync`. Every `X` row, and every `D` row that "
        "changes a commercial or contractual position. **Priority 1 = blocks submission**; "
        "Priority is the one hand-kept column and survives regeneration (default 2).\n\n"
        f"{deadline}\n**Updated:** {dt.date.today().isoformat()}\n\n"
        f"## Decisions to approve\n\n{_md_table(DECISION_COLS, decisions)}\n\n"
        f"## Firm inputs owed\n\n{_md_table(FIRM_COLS, firm)}\n\n"
        "## All commitments\n\nEvery register row, so the bid lead can review each decision — "
        "not only the ones above.\n\n"
        f"{_md_table(ALL_COLS, every)}\n", encoding="utf-8")
    print(f"signoff: wrote {md_path.relative_to(root)} · {len(decisions)} decision(s) · "
          f"{len(firm)} firm input(s) · {len(every)} commitment(s)")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        print("  signoff_sheet.xlsx NOT written — openpyxl missing (pip install openpyxl)")
        return 1
    wb = Workbook()
    wb.remove(wb.active)
    for title, cols, rows in (("Decisions to approve", DECISION_COLS, decisions),
                              ("Firm inputs owed", FIRM_COLS, firm),
                              ("All commitments", ALL_COLS, every)):
        ws = wb.create_sheet(title)
        ws.append(cols + REVIEW_COLS)
        for r in rows:
            r = list(r)
            if cols is FIRM_COLS and str(r[1]).isdigit():
                r[1] = int(r[1])
            ws.append(r + ["", "", ""])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=HEAD_FILL)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for n, col in enumerate(cols + REVIEW_COLS, 1):
            wide = col in ("Value", "Value (fallback if not confirmed)", "Basis", "New value",
                           "Comment")
            ws.column_dimensions[ws.cell(row=1, column=n).column_letter].width = 40 if wide else 16
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(xlsx_path)
    print(f"  wrote {xlsx_path.relative_to(root)} — send it to the approvers; read it back "
          "with `import-xlsx`")
    return 0


# ── import-xlsx ───────────────────────────────────────────────────────────────

def _cell(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _norm(answer):
    ans, new, comment = answer
    return ans.upper()[:1], new, comment


def _show(answer):
    return "(" + ", ".join(f"{k}={v}" for k, v in zip(("approved", "new value", "comment"),
                                                     answer) if v) + ")"


def cmd_import(reg, root, xlsx):
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("import-xlsx needs openpyxl (pip install openpyxl)")
        return 2
    by_id = {r["id"]: r for r in reg.rows}
    wb = load_workbook(xlsx, data_only=True)
    print(f"import-xlsx: {xlsx}")
    answers = collections.OrderedDict()           # ID → (answer tuple, sheet) — first wins
    for ws in wb.worksheets:
        head = [_cell(c.value) for c in ws[1]]
        if "ID" not in head or not set(REVIEW_COLS) <= set(head):
            continue
        idx = {k: head.index(k) for k in ["ID"] + REVIEW_COLS}
        for cells in ws.iter_rows(min_row=2, values_only=True):
            cid = _cell(cells[idx["ID"]]) if idx["ID"] < len(cells) else ""
            if not cid:
                continue
            got = tuple(_cell(cells[idx[k]]) if idx[k] < len(cells) else ""
                        for k in REVIEW_COLS)
            if not any(got):
                continue
            if cid not in answers:
                answers[cid] = (got, ws.title)
            elif _norm(answers[cid][0]) != _norm(got):
                print(f"  conflict  {cid} · \"{answers[cid][1]}\" says {_show(answers[cid][0])}, "
                      f"\"{ws.title}\" says {_show(got)} — the first is used")
    seen = 0
    for cid, ((ans, new, comment), _sheet) in answers.items():
        row = by_id.get(cid)
        if row is None:
            print(f"  unknown   {cid} · not in the register — skipped")
            continue
        upd = {}
        if new and new != row["value"]:
            upd = {"value": new, "status": "changed"}
            print(f"  changed   {cid} · Value \"{row['value']}\" → \"{new}\"")
        elif ans.upper() in ("Y", "YES"):
            upd = {"status": "approved"}
            print(f"  approved  {cid}")
        elif ans.upper() in ("N", "NO"):
            upd = {"status": "rejected"}
            print(f"  rejected  {cid} · no new value given — the register needs one")
        tag = f"[reviewer: {comment}]"
        if comment and tag not in row.get("basis", ""):
            upd["basis"] = (row.get("basis", "") + " " + tag).strip()
        if upd:
            reg.set(row, **upd)
            seen += 1
    reg.save()
    print(f"  {seen} row(s) updated in {reg.path.relative_to(root)}")
    lines, _ = apply(reg, root, dry_run=True)
    print("apply would do:")
    print("\n".join(lines) if lines else "  nothing — no Value changed")
    return 0


# ── rebuild ───────────────────────────────────────────────────────────────────

def _tail(text, n=12):
    return "\n".join("    " + line for line in text.strip().splitlines()[-n:])


def cmd_rebuild(reg, root):
    draft = reg.pursuit / "3_drafting"
    render = draft / "_render"
    render.mkdir(parents=True, exist_ok=True)
    rc_all = 0
    sec_root = draft / "sections"
    dirs = sorted({p.parent for p in sec_root.rglob("*.md")}) if sec_root.exists() else []
    for d in dirs:                       # render_document reads one directory per volume
        name = reg.pursuit.name if d == sec_root else f"{reg.pursuit.name}_{d.name}"
        p = subprocess.run([sys.executable, str(SCRIPTS / "render_document.py"),
                            "--sections", str(d), "--out", str(render), "--name", name,
                            "--to", "deck-manifest", "--profile", "draft"],
                           capture_output=True, text=True)
        ok = p.returncode == 0
        rc_all |= not ok
        print(f"render: {'ok' if ok else f'FAILED (exit {p.returncode})'} · "
              f"{(render / (name + '.deck-manifest.json')).relative_to(root)}")
        if not ok:
            print(_tail(p.stdout + p.stderr))
    if not dirs:
        print("render: no sections to render")
    script = render / "rebuild.sh"
    if script.exists():
        p = subprocess.run(["bash", str(script), str(root)], cwd=render,
                           capture_output=True, text=True)
        rc_all |= p.returncode != 0
        print(f"rebuild: {'ok' if p.returncode == 0 else f'FAILED (exit {p.returncode})'} · "
              f"{script.relative_to(root)}")
        print(_tail(p.stdout + p.stderr))
    else:
        print("rebuild: no rebuild.sh in 3_drafting/_render/ — deck not rebuilt "
              "(rebuild it by hand, or add the repo's deck build there)")
    decks = sorted((d for d in render.glob("*.pptx") if not d.name.startswith("~$")),
                   key=lambda d: d.stat().st_mtime)
    if decks:
        p = subprocess.run([sys.executable, str(SCRIPTS / "verify_deck.py"), str(decks[-1]),
                            "--figures", "native"], capture_output=True, text=True)
        rc_all |= p.returncode != 0
        print(f"verify: {'clean' if p.returncode == 0 else f'FAILED (exit {p.returncode})'} · "
              f"{decks[-1].relative_to(root)}")
        print(_tail(p.stdout + p.stderr))
    else:
        print("verify: no .pptx in 3_drafting/_render/ — nothing to verify")
    if report_deck_drift(reg, root):
        return 3
    return 1 if rc_all else 0


# ── main ──────────────────────────────────────────────────────────────────────

def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("root", help="engagement repo root")
    ap.add_argument("command", nargs="?", default="check",
                    choices=["check", "apply", "signoff", "import-xlsx", "rebuild", "all"])
    ap.add_argument("xlsx", nargs="?", help="the returned sign-off workbook (import-xlsx)")
    ap.add_argument("--eng-id", help="pursuit id under 01_pursuit/ (needed when there are several)")
    ap.add_argument("--strict", action="store_true", help="check: exit 1 on drift or unsynced rows")
    args = ap.parse_args()

    root = pathlib.Path(args.root).resolve()
    paths = register_paths(root, args.eng_id)
    if not paths:
        print(f"no 01_pursuit/{args.eng_id or '*'}/{REGISTER_REL} under {root} — nothing to sync")
        return 2
    if args.command != "check" and len(paths) > 1:
        print("several registers found — pass --eng-id: "
              + ", ".join(p.parents[1].name for p in paths))
        return 2
    try:
        regs = [Register(p) for p in paths]
    except ValueError as exc:
        print(exc)
        return 2

    reg = regs[0]
    if args.command == "check":
        return max(cmd_check(r, root, args.strict) for r in regs)
    if args.command == "apply":
        return cmd_apply(reg, root)
    if args.command == "signoff":
        return cmd_signoff(reg, root)
    if args.command == "import-xlsx":
        if not args.xlsx or not pathlib.Path(args.xlsx).exists():
            print("import-xlsx needs the path of the returned workbook")
            return 2
        return cmd_import(reg, root, pathlib.Path(args.xlsx))
    if args.command == "rebuild":
        return cmd_rebuild(reg, root)
    rcs = [cmd_apply(reg, root), cmd_signoff(Register(reg.path), root),
           cmd_rebuild(Register(reg.path), root),
           cmd_check(Register(reg.path), root, strict=True)]
    return 3 if 3 in rcs else (1 if any(rcs) else 0)


if __name__ == "__main__":
    sys.exit(main())
