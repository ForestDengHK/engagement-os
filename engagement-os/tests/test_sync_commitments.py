#!/usr/bin/env python3
"""Tests for sync_commitments.py — one register drives the sections, sign-off sheet and deck.

Each test builds a small pursuit repo (two sections, a commitments register, an optional fake
`_render/rebuild.sh` that writes a stdlib-zipfile .pptx like test_verify_deck.py does) and runs
the script as the user would. Runs under pytest (`python3 -m pytest tests -q`) and as a script.
"""
import importlib.util
import pathlib
import re
import subprocess
import sys
import tempfile

PLUGIN = pathlib.Path(__file__).resolve().parents[1]
SCRIPTS = PLUGIN / "skills/eng-os/scripts"
SYNC = SCRIPTS / "sync_commitments.py"
LINT = SCRIPTS / "eng_lint.py"


def _load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


sc = _load("section_contract", SCRIPTS / "section_contract.py")

ENG = "01_pursuit/26-011"
DRAFT = f"{ENG}/3_drafting"
SECS = f"{DRAFT}/sections/v1"

SEC_A = """---
section: "Commercial terms — 24 months price hold"
status: reviewed-r2
depends_on: [commitments.md]
---

# Commercial terms

Rates are held for 24 months from contract start. Hypercare runs four weeks after go-live.

The 24 months hold covers every grade.

## Review log

| Round | Reviewer / lens | Date | Verdict | What changed |
|---|---|---|---|---|
| R1 | panel red-team | 2026-10-06 | pass | proposed 24 months |
| R2 | partner | 2026-10-07 | pass | none |
"""

SEC_B = """---
section: "Rate card"
status: draft
depends_on: [commitments.md]
---

# Rate card

All day rates are fixed for 24 months. We are certified to ISO/IEC 27001.

## Review log

| Round | Reviewer / lens | Date | Verdict | What changed |
|---|---|---|---|---|
| R1 | panel red-team | <date> | <pass/revise/blocked> | <what it caught> |
"""

HEAD = ("| ID | Topic | Class | Value | Applied value | Basis | Approver | Status | Sections | Slides |\n"
        "|---|---|---|---|---|---|---|---|---|---|\n")


def row(i, topic, cls, value, applied, status, sections, basis="estimate", approver="Morgan",
        slides=""):
    return (f"| {i} | {topic} | {cls} | {value} | {applied} | {basis} | {approver} | {status} "
            f"| {sections} | {slides} |\n")


CLEAN_ROWS = [
    row("C-01", "Price hold", "D", "24 months", "24 months", "pending",
        "10_commercial.md, 17_rate_card.md", slides="4"),
    row("C-02", "Hypercare period", "D", "four weeks", "four weeks", "decided",
        "10_commercial.md", approver="Bid lead"),
    row("C-03", "ISO 27001 certificate held", "X", "certified to ISO/IEC 27001", "", "pending",
        "17_rate_card.md", basis="certificate copy from Risk", approver="Risk & Legal"),
]

RENDER_DECK = r'''#!/bin/bash
set -e
python3 - "$1" <<'EOF'
import sys, zipfile
sys.path.insert(0, "@SCRIPTS@")
import commitments
commitments.load(sys.argv[1])
hold = commitments.value("C-01")
p = ('<p:presentation xmlns:p="p" xmlns:a="a"><p:sldIdLst/>'
     '<p:sldSz cx="12192000" cy="6858000"/></p:presentation>')
s = ('<p:sld xmlns:p="p" xmlns:a="a"><p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r>'
     '<a:rPr><a:latin typeface="Arial"/></a:rPr><a:t>Rates held ' + hold + '</a:t>'
     '</a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld></p:sld>')
with zipfile.ZipFile("deck_v0.4.pptx", "w") as z:
    z.writestr("ppt/presentation.xml", p)
    z.writestr("ppt/slides/slide1.xml", s)
EOF
echo "rebuilt deck"
'''.replace("@SCRIPTS@", str(SCRIPTS))


def repo(rows=None, extra=None):
    root = pathlib.Path(tempfile.mkdtemp(prefix="engos-sync-"))
    files = {
        "CLAUDE.md": "# test\n",
        f"{SECS}/10_commercial.md": SEC_A,
        f"{SECS}/17_rate_card.md": SEC_B,
        f"{DRAFT}/commitments.md": ("# Commitments register\n\n## Register\n\n" + HEAD
                                    + "".join(rows if rows is not None else CLEAN_ROWS)),
    }
    files.update(extra or {})
    for rel, text in files.items():
        p = root / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(text, encoding="utf-8")
    return root


def run(root, *args):
    proc = subprocess.run([sys.executable, str(SYNC), str(root), *args],
                          capture_output=True, text=True)
    return proc.stdout + proc.stderr, proc.returncode


def register_rows(root):
    out = {}
    for line in (root / DRAFT / "commitments.md").read_text().splitlines():
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if cells and re.fullmatch(r"C-\d+", cells[0]):
            out[cells[0]] = cells
    return out          # cells: ID Topic Class Value Applied Basis Approver Status Sections Slides


def text(root, name):
    return (root / SECS / name).read_text(encoding="utf-8")


# ── check ─────────────────────────────────────────────────────────────────────

def test_check_clean_register():
    root = repo()
    out, rc = run(root, "check", "--strict")
    assert rc == 0, out
    assert "decided 1" in out and "pending 2" in out, out


def test_check_accepts_semicolon_separated_sections():
    rows = [row("C-01", "Price hold", "D", "24 months", "24 months", "decided",
                "10_commercial.md; 17_rate_card.md")]
    out, rc = run(repo(rows), "check", "--strict")
    assert rc == 0, out


def test_check_reports_drift_and_unsynced():
    rows = list(CLEAN_ROWS)
    rows[1] = row("C-02", "Hypercare period", "D", "six weeks", "six weeks", "decided",
                  "10_commercial.md")
    rows.append(row("C-04", "Indexation", "D", "CPI capped at 3%", "CPI capped at 4%", "changed",
                    "17_rate_card.md"))
    root = repo(rows)
    out, rc = run(root, "check")
    assert rc == 0, out                                    # informational without --strict
    assert "C-02 · 10_commercial.md · applied value not found" in out, out
    assert re.search(r"unsynced\s+C-04", out), out
    out, rc = run(root, "check", "--strict")
    assert rc == 1, out


def test_check_ignores_frontmatter_and_review_log():
    # "proposed 24 months" lives in the review log and the frontmatter title only
    rows = [row("C-05", "Proposal wording", "D", "proposed 24 months", "", "decided",
                "10_commercial.md")]
    root = repo(rows)
    out, rc = run(root, "check", "--strict")
    assert "C-05 · 10_commercial.md" in out and rc == 1, out


def test_check_flags_rejected_without_new_value():
    rows = list(CLEAN_ROWS)
    rows[0] = row("C-01", "Price hold", "D", "24 months", "24 months", "rejected",
                  "10_commercial.md, 17_rate_card.md")
    root = repo(rows)
    out, _ = run(root, "check")
    assert re.search(r"C-01.*rejected", out), out


# ── apply ─────────────────────────────────────────────────────────────────────

def test_apply_rewrites_body_reopens_and_logs():
    rows = list(CLEAN_ROWS)
    rows[0] = row("C-01", "Price hold", "D", "12 months", "24 months", "changed",
                  "10_commercial.md, 17_rate_card.md")
    root = repo(rows)
    out, rc = run(root, "apply")
    assert rc == 0, out
    assert re.search(r"10_commercial\.md:\d+", out) and re.search(r"17_rate_card\.md:\d+", out), out

    a, b = text(root, "10_commercial.md"), text(root, "17_rate_card.md")
    meta_a, body_a = sc.parse_frontmatter(a)
    log_a = body_a.split("## Review log")[1]
    assert "held for 12 months" in a and "The 12 months hold" in a, a
    assert "24 months price hold" in meta_a["section"], "frontmatter must not be rewritten"
    assert "proposed 24 months" in log_a, "existing review history must not be rewritten"
    assert meta_a["status"] == "revise-r1", meta_a
    assert sc.latest_verdict(a)[1] == "revise", a
    assert re.search(r'\| R1 \(sync\) \|.*\| revise \|.*C-01 changed: "24 months" → "12 months"', a), a

    meta_b, _ = sc.parse_frontmatter(b)
    assert "fixed for 12 months" in b and meta_b["status"] == "draft", b
    assert "(sync)" not in b, "a draft section has no review round to reopen"

    r = register_rows(root)["C-01"]
    assert r[4] == "12 months" and r[7] == "decided", r


def test_apply_is_idempotent_and_records_first_sync():
    root = repo()
    run(root, "apply")
    assert register_rows(root)["C-03"][4] == "certified to ISO/IEC 27001"   # empty → recorded
    before = text(root, "10_commercial.md")
    out, rc = run(root, "apply")
    assert rc == 0 and text(root, "10_commercial.md") == before, out
    assert "nothing to apply" in out, out


def test_apply_reports_unfound_value_and_keeps_row_unsynced():
    rows = [row("C-01", "Price hold", "D", "12 months", "36 months", "changed", "10_commercial.md")]
    root = repo(rows)
    out, rc = run(root, "apply")
    assert "C-01 · 10_commercial.md · applied value not found" in out, out
    assert register_rows(root)["C-01"][4] == "36 months" and rc == 1


# ── signoff ───────────────────────────────────────────────────────────────────

def test_signoff_generates_md_and_xlsx():
    from openpyxl import load_workbook
    prior = ("## Firm inputs owed\n\n| ID | Priority | Topic |\n|---|---|---|\n"
             "| C-03 | 1 | ISO 27001 certificate held |\n")
    root = repo(extra={f"{DRAFT}/signoff_sheet.md": prior})
    out, rc = run(root, "signoff")
    assert rc == 0, out
    md = (root / DRAFT / "signoff_sheet.md").read_text()
    decisions, firm = md.split("## Firm inputs owed")
    assert "C-01" in decisions and "C-02" not in decisions, md   # hypercare: not a sign-off topic
    assert re.search(r"\| C-03 \| 1 \|", firm), "a priority set by hand survives regeneration"

    wb = load_workbook(root / DRAFT / "signoff_sheet.xlsx")
    assert wb.sheetnames == ["Decisions to approve", "Firm inputs owed", "All commitments"], \
        wb.sheetnames
    for ws in wb.worksheets:
        head = [c.value for c in ws[1]]
        assert {"ID", "Approved? (Y/N)", "New value", "Comment"} <= set(head), head
        assert ws.freeze_panes == "A2" and ws.auto_filter.ref, ws.title
        assert ws["A1"].font.bold and ws["A1"].fill.fgColor.rgb.endswith("86BC25")
        assert ws["A1"].font.color.rgb.endswith("FFFFFF")
    ids = [ws.cell(row=2, column=1).value for ws in wb.worksheets]
    assert ids == ["C-01", "C-03", "C-01"], ids
    firm_ws = wb["Firm inputs owed"]
    assert firm_ws.cell(row=2, column=[c.value for c in firm_ws[1]].index("Priority") + 1).value == 1


def test_signoff_all_commitments_lists_every_row():
    from openpyxl import load_workbook
    root = repo()
    run(root, "signoff")
    md = (root / DRAFT / "signoff_sheet.md").read_text()
    allpart = md.split("## All commitments")[1]
    assert all(f"| {i} |" in allpart for i in ("C-01", "C-02", "C-03")), md
    assert "| ID | Topic | Class | Value | Applied value | Basis | Approver | Status | Sections " \
           "| Slides |" in allpart, allpart
    ws = load_workbook(root / DRAFT / "signoff_sheet.xlsx")["All commitments"]
    head = [c.value for c in ws[1]]
    assert head == ["ID", "Topic", "Class", "Value", "Applied value", "Basis", "Approver",
                    "Status", "Sections", "Slides", "Approved? (Y/N)", "New value", "Comment"], head
    assert [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)] == \
        ["C-01", "C-02", "C-03"]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref and ws["A1"].font.bold
    assert ws["A1"].fill.fgColor.rgb.endswith("86BC25")


def test_import_reads_any_sheet_first_answer_wins_conflicts_reported():
    root = repo()
    run(root, "signoff")
    xlsx = root / DRAFT / "signoff_sheet.xlsx"
    fill_xlsx(xlsx, {"C-02": ("N", "six weeks", None)}, sheet="All commitments")   # only there
    fill_xlsx(xlsx, {"C-01": ("Y", None, None)}, sheet="Decisions to approve")
    fill_xlsx(xlsx, {"C-01": ("N", "12 months", None)}, sheet="All commitments")  # conflicts
    out, rc = run(root, "import-xlsx", str(xlsx))
    assert rc == 0, out
    r = register_rows(root)
    assert r["C-02"][3] == "six weeks" and r["C-02"][7] == "changed", r["C-02"]
    assert r["C-01"][3] == "24 months" and r["C-01"][7] == "approved", r["C-01"]
    assert re.search(r"conflict\s+C-01", out), out


def fill_xlsx(path, answers, sheet=None):
    """answers: {ID: (approved, new value, comment)} written into the reviewer columns
    (of every sheet holding the ID, or only `sheet`)."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    for ws in wb.worksheets:
        if sheet and ws.title != sheet:
            continue
        head = [c.value for c in ws[1]]
        for r in range(2, ws.max_row + 1):
            rid = ws.cell(row=r, column=1).value
            if rid in answers:
                for col, val in zip(("Approved? (Y/N)", "New value", "Comment"), answers[rid]):
                    ws.cell(row=r, column=head.index(col) + 1).value = val
    wb.save(path)


# ── import-xlsx ───────────────────────────────────────────────────────────────

def test_import_xlsx_sets_status_value_and_comment():
    rows = list(CLEAN_ROWS) + [row("C-04", "Liability cap", "D", "1x fees", "1x fees", "pending",
                                    "10_commercial.md")]
    root = repo(rows)
    run(root, "signoff")
    xlsx = root / DRAFT / "signoff_sheet.xlsx"
    fill_xlsx(xlsx, {"C-01": ("N", "12 months", "24 is too long"),
                     "C-03": ("Y", None, None),
                     "C-04": ("N", None, None)})
    before = text(root, "10_commercial.md")
    out, rc = run(root, "import-xlsx", str(xlsx))
    assert rc == 0, out
    r = register_rows(root)
    assert r["C-01"][3] == "12 months" and r["C-01"][7] == "changed", r["C-01"]
    assert "[reviewer: 24 is too long]" in r["C-01"][5], r["C-01"]
    assert r["C-03"][7] == "approved" and r["C-04"][7] == "rejected", r
    assert re.search(r"would .*10_commercial\.md", out), out           # previews apply
    assert text(root, "10_commercial.md") == before, "import must not touch sections"


# ── rebuild ───────────────────────────────────────────────────────────────────

def test_rebuild_renders_runs_script_and_verifies_newest_deck():
    root = repo(extra={f"{DRAFT}/_render/rebuild.sh": RENDER_DECK})
    out, rc = run(root, "rebuild")
    assert rc == 0, out
    assert list((root / DRAFT / "_render").glob("*.deck-manifest.json")), out
    assert "rebuilt deck" in out and "deck_v0.4.pptx" in out, out
    assert re.search(r"verify.*clean|✓ clean", out), out


def test_rebuild_without_script_or_deck_says_so():
    root = repo()
    out, rc = run(root, "rebuild")
    assert rc == 0, out
    assert "no rebuild.sh" in out and "no .pptx" in out, out


# ── all (the Morgan scenario) ─────────────────────────────────────────────────

def test_all_after_import_leaves_everything_clean():
    root = repo(extra={f"{DRAFT}/_render/rebuild.sh": RENDER_DECK})
    run(root, "signoff")
    xlsx = root / DRAFT / "signoff_sheet.xlsx"
    fill_xlsx(xlsx, {"C-01": ("N", "12 months", None)})
    run(root, "import-xlsx", str(xlsx))
    out, rc = run(root, "all")
    assert rc == 0, out
    assert "fixed for 12 months" in text(root, "17_rate_card.md")
    assert sc.parse_frontmatter(text(root, "10_commercial.md"))[0]["status"] == "revise-r1"
    md = (root / DRAFT / "signoff_sheet.md").read_text()
    assert "12 months" in md, md
    out, rc = run(root, "check", "--strict")
    assert rc == 0, out


# ── commitments.py helper ─────────────────────────────────────────────────────

def test_helper_value_and_errors():
    cm = _load("commitments", SCRIPTS / "commitments.py")
    rows = list(CLEAN_ROWS) + [
        row("C-07", "Rate hold", "D", "fixed until 31 December 2027", "", "decided", ""),
        row("C-08", "Start date", "D", "starting 2027-01-04", "", "decided", ""),
        row("C-09", "Owed", "X", "", "", "pending", ""),
        row("C-10", "Hypercare", "D", "four weeks", "", "decided", "")]
    root = repo(rows)
    regs = cm.load(root)                                   # repo root …
    assert regs["C-01"]["value"] == "24 months", regs["C-01"]
    assert cm.load(root / DRAFT / "commitments.md")["C-02"]["value"] == "four weeks"  # … or path
    assert cm.value("C-07") == "fixed until 31 December 2027"
    import datetime
    assert cm.date_of("C-07") == datetime.date(2027, 12, 31)
    assert cm.date_of("C-08") == datetime.date(2027, 1, 4)
    for bad, exc in (("C-99", KeyError), ("C-09", ValueError)):
        try:
            cm.value(bad)
        except exc:
            pass
        else:
            raise AssertionError(f"{bad} should raise {exc.__name__}")
    assert cm.value("C-09", default="to be confirmed") == "to be confirmed"
    try:
        cm.date_of("C-10")
    except ValueError:
        pass
    else:
        raise AssertionError("a Value with no date must raise ValueError")


# ── deck drift ────────────────────────────────────────────────────────────────

def literal_deck(path, *texts):
    """A stdlib-zipfile .pptx (as test_verify_deck.py builds) with one textbox per slide.
    The first slide splits its text across two runs, as PowerPoint does."""
    import zipfile
    p = ('<p:presentation xmlns:p="p" xmlns:a="a"><p:sldIdLst/>'
         '<p:sldSz cx="12192000" cy="6858000"/></p:presentation>')
    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("ppt/presentation.xml", p)
        for i, t in enumerate(texts, 1):
            head, _, tail = t.partition(" ")
            runs = (f'<a:r><a:t>{head} </a:t></a:r><a:r><a:t>{tail}</a:t></a:r>' if i == 1
                    else f'<a:r><a:t>{t}</a:t></a:r>')
            z.writestr(f"ppt/slides/slide{i}.xml",
                       '<p:sld xmlns:p="p" xmlns:a="a"><p:cSld><p:spTree><p:sp><p:txBody>'
                       f'<a:p>{runs}</a:p></p:txBody></p:sp></p:spTree></p:cSld></p:sld>')


def test_check_deck_drift_reports_miss_and_exits_3():
    root = repo()                                     # C-01 Slides=4, Applied=24 months
    literal_deck(root / DRAFT / "_render/old.pptx", "Rates held 24 months")
    literal_deck(root / DRAFT / "_render/deck_v2.pptx", "Rates held 36 months")   # newest
    out, rc = run(root, "check")
    assert rc == 3, out
    assert "MISS C-01: '24 months' not in deck (slides 4)" in out, out
    assert "deck_v2.pptx" in out, "the newest deck is the one checked"
    _, rc = run(root, "check", "--strict")
    assert rc == 3


def test_check_deck_drift_clean_when_value_on_slide():
    root = repo()
    literal_deck(root / DRAFT / "_render/deck.pptx", "Rates are held 24 months")  # split runs
    out, rc = run(root, "check", "--strict")
    assert rc == 0 and "MISS" not in out, out
    assert re.search(r"deck-drift.*0 miss", out), out


def test_rebuild_reports_deck_drift_after_verify():
    stale = RENDER_DECK.replace('hold = commitments.value("C-01")', 'hold = "36 months"')
    root = repo(extra={f"{DRAFT}/_render/rebuild.sh": stale})
    out, rc = run(root, "rebuild")
    assert rc == 3 and "MISS C-01" in out, out
    assert out.index("verify:") < out.index("deck-drift"), out


# ── lint hook ─────────────────────────────────────────────────────────────────

def lint(root):
    proc = subprocess.run([sys.executable, str(LINT), str(root)], capture_output=True, text=True)
    return proc.stdout


def test_lint_commitment_drift_and_rejected_only_at_freeze():
    rows = list(CLEAN_ROWS)
    rows[0] = row("C-01", "Price hold", "D", "24 months", "24 months", "rejected",
                  "10_commercial.md")
    rows[1] = row("C-02", "Hypercare period", "D", "six weeks", "six weeks", "decided",
                  "10_commercial.md")
    root = repo(rows)
    out = lint(root)
    assert re.search(r"warn\s+\[commitment-drift\]", out), out
    assert "commitment-rejected-unresolved" not in out, out
    (root / ENG / "4_final").mkdir(parents=True)
    (root / ENG / "4_final/response.docx").write_text("x")
    out = lint(root)
    assert re.search(r"ERROR\s+\[commitment-rejected-unresolved\]", out), out


def test_lint_reports_score_gaps_as_info():
    gaps = ("# Score-gap register\n\n"
            "| ID | Round | Criterion | Score | Gap (what the evaluator saw) | Improvement "
            "| Type (firm input / rewrite) | Owner | Sections/slides "
            "| Status (open / in progress / applied / won't fix) |\n"
            "|---|---|---|---|---|---|---|---|---|---|\n"
            "| G-01 | R1b | Approach | 5 | generic | name the method | rewrite | lead | 10 | open |\n"
            "| G-02 | R1b | Approach | 5 | no proof | client reference | firm input | partner | 10 | open |\n"
            "| G-03 | R1b | Price | 6 | vague | table | rewrite | lead | 17 | applied |\n")
    root = repo(extra={f"{ENG}/2_analysis/score_gap_register.md": gaps})
    out = lint(root)
    assert re.search(r"info\s+\[score-gaps\]", out), out
    assert "1 rewrite gap(s) blocking the next evaluator round" in out, out
    assert "1 firm-input gap(s) owed" in out, out


def main():
    tests = [(n, f) for n, f in list(globals().items()) if n.startswith("test_") and callable(f)]
    fails = []
    for name, fn in tests:
        try:
            fn()
            print(f"  ✓ {name}")
        except Exception as exc:                     # a crash is a failure, not an abort
            fails.append(name)
            print(f"  ✗ {name}\n      {type(exc).__name__}: {str(exc)[:600]}")
    print(f"\n{len(tests) - len(fails)}/{len(tests)} sync-commitments cases pass")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
