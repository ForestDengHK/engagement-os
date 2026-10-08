#!/usr/bin/env python3
"""Regression tests for build_estimate_workbook.py — reasoning carried, parameters not constants,
and the panel / rate-card shape.

Each case pins a defect seen on a live tender: the builder dropped every reasoning section of the
seed (the workbook then said "8 of 8 owed" and the export wrote that over the seed), and it
hardcoded €, a 26-week term and 300 marks on a USD rate-card panel with no price formula.
"""
import importlib.util
import pathlib
import subprocess
import sys

import openpyxl
import pytest

PLUGIN = pathlib.Path(__file__).resolve().parents[1]
BUILDER = PLUGIN / "skills/eng-os/scripts/build_estimate_workbook.py"

spec = importlib.util.spec_from_file_location("build_estimate_workbook", BUILDER)
bw = importlib.util.module_from_spec(spec)
sys.path.insert(0, str(BUILDER.parent))
spec.loader.exec_module(bw)

EFFORT = """<!--table:effort-->
| S-ID | Activity | Driver | O | M | P | Expected | Lead |
|---|---|---|---|---|---|---|---|
| S-01 | Assess | 3 x 5d | 15 | 20 | 28 | 20 | Lead A |
| S-02 | Report | 2 x 8d | 16 | 22 | 30 | 22 | Lead B |
"""

SECTIONS = {
    "BasisOfEstimate": ("## 1. Basis of estimate",
                        "Sized from the buyer's own volumetrics: 18 loads, 3 environments."),
    "Calibration": ("## 6. Calibration & sanity checks",
                    "Two comparable assessments ran 160 and 190 days; this sits between them."),
    "Triggers": ("## 11. Re-baseline triggers",
                 "Re-baseline if the estate exceeds 25 source systems or the term moves."),
}


def seed(extra="", sections=SECTIONS):
    body = "# Estimation\n\n" + extra + "\n" + EFFORT
    for head, text in sections.values():
        body += f"\n{head}\n\n{text}\n"
    return body


def run(*args):
    return subprocess.run([sys.executable, str(BUILDER), *map(str, args)],
                          capture_output=True, text=True)


def every_value(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                yield c


def judgement_rows(wb):
    ws = wb[bw.JUDGEMENT_SHEET]
    return {r[0].value: r[2].value for r in ws.iter_rows(min_row=2) if r[0].value
            and r[0].value != "JUDGEMENT OWED"}


def name_value(wb, name):
    sheet, ref = next(iter(wb.defined_names[name].destinations))
    return wb[sheet][ref.replace("$", "")].value


# ── A1. reasoning sections ─────────────────────────────────────────────────────

def test_seed_reasoning_sections_land_in_judgement_and_seed_unchanged(tmp_path):
    md = tmp_path / "estimation.md"
    md.write_text(seed(), encoding="utf-8")
    before = md.read_bytes()
    r = run(md, "--out", tmp_path / "est.xlsx", "--no-recalc")
    assert r.returncode == 0, r.stderr
    rows = judgement_rows(openpyxl.load_workbook(tmp_path / "est.xlsx"))
    written = {k: v for k, v in rows.items() if v != "<not yet written>"}
    assert set(written) == set(SECTIONS)
    for topic, (_h, text) in SECTIONS.items():
        assert text in written[topic]
    assert md.read_bytes() == before


def test_seed_heading_case_does_not_drop_a_section(tmp_path):
    sections = {"BasisOfEstimate": ("## 1. Basis Of Estimate", "Bottom-up from 14 interviews.")}
    md = tmp_path / "estimation.md"
    md.write_text(seed(sections=sections), encoding="utf-8")
    assert run(md, "--out", tmp_path / "est.xlsx", "--no-recalc").returncode == 0
    rows = judgement_rows(openpyxl.load_workbook(tmp_path / "est.xlsx"))
    assert "Bottom-up from 14 interviews." in rows["BasisOfEstimate"]


def test_export_refuses_to_write_owed_over_a_seed_section_with_content(tmp_path):
    # the field failure: a workbook built without the seed's reasoning, exported over the seed
    md = tmp_path / "estimation.md"
    md.write_text(seed(), encoding="utf-8")
    before = md.read_bytes()
    xlsx = tmp_path / "estimation.xlsx"
    bw.build(bw.parse_tables(seed()), 0.5, str(xlsx))          # no narrative: every row OWED
    r = run(md, "--out", xlsx, "--to-md", "--no-recalc")
    assert r.returncode == 6, r.stdout + r.stderr
    assert "BasisOfEstimate" in r.stderr
    assert md.read_bytes() == before


# ── A2. parameters, not constants ──────────────────────────────────────────────

def test_defaults_are_unchanged(tmp_path):
    md = tmp_path / "estimation.md"
    md.write_text(seed(), encoding="utf-8")
    assert run(md, "--out", tmp_path / "e.xlsx", "--no-recalc").returncode == 0
    wb = openpyxl.load_workbook(tmp_path / "e.xlsx")
    assert name_value(wb, "term_weeks") == 26
    assert name_value(wb, "max_marks") == 300
    assert name_value(wb, "term_days") == 126
    assert wb["Gantt"].max_column == 27
    assert wb["Cost"]["B1"].value == "at P50 €"


def test_currency_flag_leaves_no_euro_anywhere(tmp_path):
    md = tmp_path / "estimation.md"
    md.write_text(seed(), encoding="utf-8")
    assert run(md, "--out", tmp_path / "e.xlsx", "--no-recalc", "--currency", "USD").returncode == 0
    wb = openpyxl.load_workbook(tmp_path / "e.xlsx")
    euro = [f"{c.parent.title}!{c.coordinate}" for c in every_value(wb)
            if "€" in str(c.value or "") or "€" in (c.number_format or "")]
    assert euro == []
    assert wb["RateCard"]["B4"].value == "Cost rate / day $"
    assert wb["Cost"]["B2"].number_format.startswith("$")


def test_term_weeks_flag_drives_term_gantt_and_default_gates(tmp_path):
    md = tmp_path / "estimation.md"
    md.write_text(seed(), encoding="utf-8")
    assert run(md, "--out", tmp_path / "e.xlsx", "--no-recalc", "--term-weeks", "14").returncode == 0
    wb = openpyxl.load_workbook(tmp_path / "e.xlsx")
    assert name_value(wb, "term_weeks") == 14
    assert wb["Gantt"].max_column == 15
    sheet, ref = next(iter(wb.defined_names["gate_weeks"].destinations))
    gates = [c[0].value for c in wb[sheet][ref.replace("$", "")]]
    assert max(gates) <= 14


def test_marks_flag_sets_the_marks_input(tmp_path):
    md = tmp_path / "estimation.md"
    md.write_text(seed(), encoding="utf-8")
    assert run(md, "--out", tmp_path / "e.xlsx", "--no-recalc", "--marks", "40").returncode == 0
    assert name_value(openpyxl.load_workbook(tmp_path / "e.xlsx"), "max_marks") == 40


def test_marks_zero_omits_the_price_to_marks_block(tmp_path):
    md = tmp_path / "estimation.md"
    md.write_text(seed(), encoding="utf-8")
    assert run(md, "--out", tmp_path / "e.xlsx", "--no-recalc", "--marks", "0").returncode == 0
    wb = openpyxl.load_workbook(tmp_path / "e.xlsx")
    assert "max_marks" not in wb.defined_names
    assert not any("max_marks" in str(c.value) or "low_a" in str(c.value) for c in every_value(wb))
    header = [c.value for c in wb["Decision"][1]]
    assert not any("Marks" in str(h) for h in header)
    assert "Candidate price" in header and "Flag" in header


def test_seed_parameters_block_and_cli_override(tmp_path):
    params = ("<!--table:parameters-->\n| Parameter | Value | Note |\n|---|---|---|\n"
              "| currency | USD | |\n| term_weeks | 14 | |\n| marks | 0 | no price formula |\n")
    md = tmp_path / "estimation.md"
    md.write_text(seed(extra=params), encoding="utf-8")
    assert run(md, "--out", tmp_path / "a.xlsx", "--no-recalc").returncode == 0
    wb = openpyxl.load_workbook(tmp_path / "a.xlsx")
    assert name_value(wb, "term_weeks") == 14
    assert "max_marks" not in wb.defined_names
    assert wb["Cost"]["B1"].value == "at P50 $"
    # a CLI flag wins over the seed
    assert run(md, "--out", tmp_path / "b.xlsx", "--no-recalc", "--currency", "GBP").returncode == 0
    assert openpyxl.load_workbook(tmp_path / "b.xlsx")["Cost"]["B1"].value == "at P50 £"


# ── A3. panel / rate-card shape ────────────────────────────────────────────────

TIERS = [("Onshore", "100%"), ("Nearshore", "70%"), ("Offshore", "45%")]
BANDS = [(f"Band {i}", 400 + 200 * i) for i in range(1, 8)]


def panel_seed():
    t = "## 0. Parameters\n\n<!--table:parameters-->\n| Parameter | Value | Note |\n|---|---|---|\n"
    t += "| shape | panel | |\n| currency | USD | |\n| term_weeks | 14 | |\n| marks | 0 | |\n"
    t += "\n<!--table:bands-->\n| Band | Base rate / day | Description |\n|---|---|---|\n"
    t += "".join(f"| {b} | ${r:,} | level {b[-1]} |\n" for b, r in BANDS)
    t += "\n<!--table:tiers-->\n| Tier | Index | Note |\n|---|---|---|\n"
    t += "".join(f"| {n} | {i} | |\n" for n, i in TIERS)
    t += "\n<!--table:roles-->\n| Role | Band | What they do |\n|---|---|---|\n"
    t += "".join(f"| Role {i} | Band {i} | does {i} |\n" for i in range(1, 8))
    t += ("\n<!--table:discounts-->\n| From spend | Discount | Basis |\n|---|---|---|\n"
          "| 0 | 0% | list |\n| 100000 | 3% | volume |\n| 250000 | 5% | volume |\n")
    t += ("\n<!--table:rom-->\n| Role | Tier | O | M | P | Note |\n|---|---|---|---|---|---|\n"
          "| Role 6 | Onshore | 8 | 10 | 15 | lead |\n| Role 4 | Nearshore | 30 | 40 | 60 | |\n"
          "| Role 2 | Offshore | 40 | 50 | 80 | |\n")
    t += ("\n<!--table:assumptions-->\n| ID | Assumption | Basis |\n|---|---|---|\n"
          "| PA1 | Remote-first delivery | RFP §4 |\n| PA2 | Rates held 12 months | firm policy |\n")
    for head, text in SECTIONS.values():
        t += f"\n{head}\n\n{text}\n"
    return "# Estimation\n\n" + t


PANEL_SHEETS = {"RateCard", "RoleCatalogue", "Discounts", "Assumptions", "WorkedROM"}


def test_panel_shape_builds_the_five_sheets(tmp_path):
    md = tmp_path / "estimation.md"
    md.write_text(panel_seed(), encoding="utf-8")
    r = run(md, "--out", tmp_path / "p.xlsx", "--no-recalc", "--ceiling-pct", "15")
    assert r.returncode == 0, r.stderr
    wb = openpyxl.load_workbook(tmp_path / "p.xlsx")
    assert PANEL_SHEETS | {"README", "Inputs", bw.JUDGEMENT_SHEET} == set(wb.sheetnames)
    # 7 bands x 3 tiers, every rate a formula over base x tier index
    sheet, ref = next(iter(wb.defined_names["pnl_grid"].destinations))
    grid = list(wb[sheet][ref.replace("$", "")])
    assert len(grid) == 7 and all(len(row) == 3 for row in grid)
    assert all(str(c.value).startswith("=") for row in grid for c in row)
    assert name_value(wb, "ceiling_pct") == 0.15
    assert {"rom_P50", "rom_P80", "rom_ceiling"} <= set(wb.defined_names)
    rows = judgement_rows(wb)
    assert all(text in rows[topic] for topic, (_h, text) in SECTIONS.items())
    assert not any("€" in str(c.value or "") or "€" in (c.number_format or "")
                   for c in every_value(wb))


def test_project_shape_has_no_panel_sheets(tmp_path):
    md = tmp_path / "estimation.md"
    md.write_text(seed(), encoding="utf-8")
    assert run(md, "--out", tmp_path / "e.xlsx", "--no-recalc").returncode == 0
    names = set(openpyxl.load_workbook(tmp_path / "e.xlsx").sheetnames)
    assert not (names & (PANEL_SHEETS - {"RateCard"}))
    assert {"Effort", "Range", "Schedule", "Gantt", "Cost", "Decision"} <= names


@pytest.mark.skipif(bw.find_recalc() is None, reason="xlsx skill recalc.py not installed")
def test_panel_workbook_recalculates_with_zero_errors(tmp_path):
    md = tmp_path / "estimation.md"
    md.write_text(panel_seed(), encoding="utf-8")
    out = tmp_path / "p.xlsx"
    r = run(md, "--out", out, "--ceiling-pct", "15")
    assert r.returncode == 0, r.stdout + r.stderr
    assert "0 errors" in r.stdout
    wb = openpyxl.load_workbook(out, data_only=True)
    p50, p80, ceiling = (name_value(wb, n) for n in ("rom_P50", "rom_P80", "rom_ceiling"))
    assert 0 < p50 < p80
    # Role 6 = Band 6 (1600) onshore x 100%; Role 4 = Band 4 (1200) x 70%; Role 2 = Band 2 (800) x 45%
    expected = (8 + 40 + 15) / 6 * 1600 + (30 + 160 + 60) / 6 * 840 + (40 + 200 + 80) / 6 * 360
    assert abs(p50 - expected) < 0.01
    net_p50 = name_value(wb, "rom_P50_net")
    assert abs(ceiling - net_p50 * 1.15) < 0.01
    # the snapshot shows money as money — LibreOffice re-saves `$#,##0` escaped as `\\$#,##0`
    snap = tmp_path / "snap.md"
    snap.write_text("x", encoding="utf-8")
    assert run(snap, "--out", out, "--to-md").returncode == 0
    text = snap.read_text(encoding="utf-8")
    assert "| $1,600 |" in text and "€" not in text
